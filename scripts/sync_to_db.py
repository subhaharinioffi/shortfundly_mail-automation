import sqlite3
import shutil
import os
import openpyxl

DB_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "database.sqlite"))
BACKUP_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "database.sqlite.bak"))
EXCEL_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "Orissa_Colleges_Complete_Email_Directory.xlsx"))

def sync_excel_to_db():
    print(f"Backing up current database to {BACKUP_PATH}...")
    shutil.copy2(DB_PATH, BACKUP_PATH)
    
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    all_leads = {}

    # 1. Load from SAMS Directory
    if 'Govt SAMS Degree Directory' in wb.sheetnames:
        ws4 = wb['Govt SAMS Degree Directory']
        for r in range(2, ws4.max_row + 1):
            cname = ws4.cell(row=r, column=4).value
            dist = ws4.cell(row=r, column=5).value
            pname = ws4.cell(row=r, column=9).value
            pmob = ws4.cell(row=r, column=10).value
            pem = ws4.cell(row=r, column=11).value
            ain = ws4.cell(row=r, column=12).value
            aimob = ws4.cell(row=r, column=13).value
            if pem and '@' in str(pem):
                e = str(pem).strip().lower()
                notes = f"Principal Mobile: {pmob}" if pmob else ""
                if aimob:
                    notes += f" | Admission In-Charge: {ain} ({aimob})"
                all_leads[e] = {
                    'state': 'Odisha',
                    'city': str(dist or 'Odisha').strip(),
                    'college_name': str(cname or '').strip(),
                    'website_url': '',
                    'department': 'Office of the Principal (Govt SAMS Verified)',
                    'contact_person': str(pname or 'Principal').strip(),
                    'email': e,
                    'source_page_url': 'https://degree.samsodisha.gov.in/',
                    'notes': notes
                }

    # 2. Load from Staff & Faculty Directory
    if 'Staff & Faculty Directory' in wb.sheetnames:
        ws3 = wb['Staff & Faculty Directory']
        for r in range(2, ws3.max_row + 1):
            cname = ws3.cell(row=r, column=2).value
            dist = ws3.cell(row=r, column=3).value
            pname = ws3.cell(row=r, column=4).value
            role = ws3.cell(row=r, column=5).value
            em = ws3.cell(row=r, column=6).value
            phone = ws3.cell(row=r, column=7).value
            src = ws3.cell(row=r, column=8).value
            if em and '@' in str(em):
                e = str(em).strip().lower()
                if e not in all_leads:
                    all_leads[e] = {
                        'state': 'Odisha',
                        'city': str(dist or 'Odisha').strip(),
                        'college_name': str(cname or '').strip(),
                        'website_url': '',
                        'department': str(role or 'Faculty / Staff').strip(),
                        'contact_person': str(pname or 'Department Official').strip(),
                        'email': e,
                        'source_page_url': str(src or '').strip(),
                        'notes': f"Direct Phone: {phone}" if phone else ""
                    }
                else:
                    if pname and all_leads[e]['contact_person'] in ('Principal', 'Department Official', ''):
                        all_leads[e]['contact_person'] = str(pname).strip()

    # 3. Load from Master Sheet for Websites & Official Emails
    if 'All Colleges Verified Directory' in wb.sheetnames:
        ws1 = wb['All Colleges Verified Directory']
        for r in range(2, ws1.max_row + 1):
            cname = ws1.cell(row=r, column=2).value
            dist = ws1.cell(row=r, column=4).value
            web = ws1.cell(row=r, column=5).value
            pname = ws1.cell(row=r, column=6).value
            pmob = ws1.cell(row=r, column=8).value
            all_em_str = ws1.cell(row=r, column=12).value
            phones = ws1.cell(row=r, column=15).value
            addr = ws1.cell(row=r, column=16).value
            src = ws1.cell(row=r, column=21).value or 'https://www.indcareer.com/'

            if all_em_str:
                for raw_em in str(all_em_str).split(','):
                    raw_em = raw_em.strip().lower()
                    if raw_em and '@' in raw_em:
                        if raw_em in all_leads:
                            if web and not all_leads[raw_em]['website_url']:
                                all_leads[raw_em]['website_url'] = str(web).strip()
                            if addr and 'Address:' not in all_leads[raw_em]['notes']:
                                all_leads[raw_em]['notes'] += f" | Address: {addr}"
                        else:
                            all_leads[raw_em] = {
                                'state': 'Odisha',
                                'city': str(dist or 'Odisha').strip(),
                                'college_name': str(cname or '').strip(),
                                'website_url': str(web or '').strip(),
                                'department': 'Administration / General',
                                'contact_person': str(pname or 'Administrative Office').strip(),
                                'email': raw_em,
                                'source_page_url': str(src).strip(),
                                'notes': f"Phones: {phones} | Address: {addr}"
                            }

    print(f"Total compiled unique verified leads from Excel: {len(all_leads)}")

    # Connect to SQLite
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    cur.execute("SELECT COUNT(*) FROM colleges;")
    initial_count = cur.fetchone()[0]
    print(f"Initial colleges count in SQLite: {initial_count}")

    # Upsert leads
    inserted = 0
    updated = 0

    for email, lead in all_leads.items():
        # Check if email exists
        cur.execute("SELECT id, contact_person, department, website_url, notes FROM colleges WHERE email = ? COLLATE NOCASE;", (email,))
        row = cur.fetchone()
        if row:
            cid, existing_person, existing_dept, existing_web, existing_notes = row
            # Update fields if new data is more descriptive
            new_person = lead['contact_person'] if (lead['contact_person'] and lead['contact_person'] not in ('Principal', 'Department Official', 'Administrative Office')) else existing_person
            new_dept = lead['department'] if lead['department'] else existing_dept
            new_web = lead['website_url'] if lead['website_url'] else existing_web
            new_notes = (existing_notes or "") + (" | " + lead['notes'] if lead['notes'] and lead['notes'] not in (existing_notes or "") else "")
            
            cur.execute("""
                UPDATE colleges 
                SET contact_person = ?, department = ?, website_url = ?, notes = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?;
            """, (new_person, new_dept, new_web, new_notes.strip(" |"), cid))
            updated += 1
        else:
            # Insert fresh lead
            cur.execute("""
                INSERT INTO colleges (state, city, college_name, website_url, department, contact_person, email, source_page_url, status, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'uncontacted', ?);
            """, (
                lead['state'],
                lead['city'],
                lead['college_name'],
                lead['website_url'],
                lead['department'],
                lead['contact_person'],
                email,
                lead['source_page_url'],
                lead['notes']
            ))
            inserted += 1

    conn.commit()
    cur.execute("SELECT COUNT(*) FROM colleges;")
    final_count = cur.fetchone()[0]
    conn.close()

    print(f"Sync Complete! Inserted: {inserted} new records, Updated: {updated} existing records.")
    print(f"New total colleges count in SQLite: {final_count}")

if __name__ == '__main__':
    sync_excel_to_db()
