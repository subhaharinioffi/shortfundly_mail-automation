import openpyxl
import requests
from bs4 import BeautifulSoup
import re
import urllib3
from urllib.parse import urljoin, urlparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import time

urllib3.disable_warnings()

headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
    'Accept-Language': 'en-US,en;q=0.5'
}

TARGET_KEYWORDS = [
    'placement', 'placements', 'tpo', 'training', 'career', 'corporate',
    'student', 'welfare', 'dsw', 'contact', 'dean', 'hod', 'iiic'
]

EMAIL_REGEX = re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}')
PHONE_REGEX = re.compile(r'(?:\+91[\-\s]?)?[6789]\d{9}|0\d{2,4}[\-\s]?\d{6,8}')

JUNK_EMAIL_DOMAINS = {
    'example.com', 'domain.com', 'email.com', 'wixpress.com', 'sentry.io',
    'wordpress.org', 'wordpress.com', 'schema.org', 'cloudflare.com', 'w3.org'
}

JUNK_EMAIL_EXTENSIONS = ('.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.css', '.js')

def is_valid_email(email):
    if not email or len(email) < 5 or len(email) > 80:
        return False
    email = email.lower().strip()
    if any(email.endswith(ext) for ext in JUNK_EMAIL_EXTENSIONS):
        return False
    parts = email.split('@')
    if len(parts) != 2:
        return False
    user, dom = parts
    if dom in JUNK_EMAIL_DOMAINS or not '.' in dom:
        return False
    return True

def crawl_site(domain_info):
    base_url = domain_info['url']
    name = domain_info['name']
    results = []

    try:
        r = requests.get(base_url, headers=headers, timeout=8, verify=False)
        if r.status_code >= 400:
            return results
        soup = BeautifulSoup(r.text, 'html.parser')
        
        # 1. Extract from homepage
        page_text = soup.get_text()
        emails = set(EMAIL_REGEX.findall(page_text))
        mailtos = set([a['href'].replace('mailto:', '').split('?')[0].strip() 
                       for a in soup.find_all('a', href=True) if a['href'].startswith('mailto:')])
        all_emails = [e for e in emails.union(mailtos) if is_valid_email(e)]
        phones = list(set(PHONE_REGEX.findall(page_text)))[:3]

        for em in all_emails:
            results.append({
                'college_name': name,
                'url': base_url,
                'email': em,
                'phone': phones[0] if phones else '',
                'source_page': base_url
            })

        # 2. Discover subpages for placement / contact / student welfare
        subpage_urls = set()
        for a in soup.find_all('a', href=True):
            href = a['href'].strip()
            text = a.get_text().lower().strip()
            if any(k in href.lower() or k in text for k in TARGET_KEYWORDS):
                full = urljoin(base_url, href)
                parsed = urlparse(full)
                if parsed.netloc.lower().replace('www.', '') == urlparse(base_url).netloc.lower().replace('www.', ''):
                    if not full.endswith(('.pdf', '.doc', '.docx', '.jpg', '.png', '#')):
                        subpage_urls.add(full)
                        if len(subpage_urls) >= 4:
                            break

        # 3. Crawl top 3 subpages
        for sub_url in list(subpage_urls)[:3]:
            try:
                sub_r = requests.get(sub_url, headers=headers, timeout=6, verify=False)
                if sub_r.status_code == 200:
                    sub_soup = BeautifulSoup(sub_r.text, 'html.parser')
                    sub_text = sub_soup.get_text()
                    sub_emails = set(EMAIL_REGEX.findall(sub_text))
                    sub_mailtos = set([a['href'].replace('mailto:', '').split('?')[0].strip() 
                                      for a in sub_soup.find_all('a', href=True) if a['href'].startswith('mailto:')])
                    sub_all_emails = [e for e in sub_emails.union(sub_mailtos) if is_valid_email(e)]
                    sub_phones = list(set(PHONE_REGEX.findall(sub_text)))[:3]

                    for em in sub_all_emails:
                        results.append({
                            'college_name': name,
                            'url': base_url,
                            'email': em,
                            'phone': sub_phones[0] if sub_phones else '',
                            'source_page': sub_url
                        })
            except Exception:
                pass

    except Exception:
        pass

    return results

if __name__ == '__main__':
    wb = openpyxl.load_workbook('C:/Users/Subhaharini/Downloads/odisha_colleges_verified.xlsx', data_only=True)
    sheet = wb.active
    domains = {}
    for r in range(4, sheet.max_row + 1):
        name = sheet.cell(row=r, column=2).value
        url = sheet.cell(row=r, column=5).value
        if url and str(url).strip().startswith('http'):
            u = str(url).strip()
            parsed = urlparse(u)
            dom = parsed.netloc.lower().replace('www.', '')
            if dom not in domains:
                domains[dom] = {'name': name, 'url': u}

    domain_list = list(domains.values())
    print(f'Starting concurrent crawl of {len(domain_list)} college/university domains with 20 threads...')
    start_time = time.time()

    all_scraped = []
    with ThreadPoolExecutor(max_workers=20) as executor:
        futures = {executor.submit(crawl_site, d): d for d in domain_list}
        completed = 0
        for f in as_completed(futures):
            res = f.result()
            if res:
                all_scraped.extend(res)
            completed += 1
            if completed % 25 == 0 or completed == len(domain_list):
                print(f'Progress: {completed}/{len(domain_list)} domains crawled ({len(all_scraped)} contacts found so far)...')

    print(f'Crawl completed in {time.time() - start_time:.2f}s! Total contacts extracted: {len(all_scraped)}')
    
    import json
    with open('scripts/crawled_contacts.json', 'w', encoding='utf-8') as f:
        json.dump(all_scraped, f, indent=2, ensure_ascii=False)
    print('Saved scraped contacts to scripts/crawled_contacts.json')

    # Show sample
    sample_placement = [c for c in all_scraped if any(k in c['email'].lower() for k in ['placement', 'tpo', 'career', 'train', 'dsw'])]
    print(f'High-priority placement/DSW emails found: {len(sample_placement)}')
    for c in sample_placement[:15]:
        print(f"  {c['college_name']} | {c['email']} | {c['phone']} | {c['source_page']}")
